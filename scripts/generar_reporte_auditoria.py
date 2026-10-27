"""
Generador de Reporte de Auditoría y Conciliación de Neumáticos
Transportes Díaz Acarreos Generales S.A.C.

Genera un libro Excel con 4 hojas profesionales:
1. RESUMEN_CONCILIACION: Explicación de la Opción B, por qué 804 y no 816, y por qué 780 únicos.
2. MENOS_DE_24_NEUMATICOS: Detalle de los 7 vehículos con menos de 24 neumáticos (12 faltantes).
3. DETALLE_DUPLICADOS: Registro de los 24 códigos duplicados (48 asignaciones).
4. FLOTA_COMPLETA_34: Estado detallado de las 34 unidades activas.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict


def crear_reporte_auditoria_excel(datos, ruta_salida):
    """
    Crea el archivo Excel de auditoría y conciliación.
    """
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    wb.remove(wb.active)

    tractos = datos['tractos']
    cisternas = datos['cisternas']
    movimientos = datos['movimientos']
    neumaticos = datos['neumaticos']

    # Estilos estándar
    font_titulo = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_subtitulo = Font(name="Calibri", size=11, italic=True, color="D9E1F2")
    font_seccion = Font(name="Calibri", size=12, bold=True, color="1F4E78")
    font_cabecera = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)
    font_alerta = Font(name="Calibri", size=10, bold=True, color="9C0006")

    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_cabecera = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_subcabecera = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_alerta = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_kpi = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_celda = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=Side(border_style="medium", color="1F4E78"), bottom=Side(border_style="medium", color="1F4E78"))

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Mapeo de unidades
    placa_to_tracto = {}
    for tr, tr_data in tractos.items():
        placa_to_tracto[tr] = tr
        if tr_data.get('cisterna'):
            placa_to_tracto[tr_data['cisterna']] = tr

    movs_por_unidad = defaultdict(lambda: {'tracto': [], 'cisterna': []})
    for m in movimientos:
        p = m['placa']
        tr = placa_to_tracto.get(p, p)
        if p in tractos:
            movs_por_unidad[tr]['tracto'].append(m)
        else:
            movs_por_unidad[tr]['cisterna'].append(m)

    # Conteo de duplicados
    codigos_conteo = defaultdict(list)
    for m in movimientos:
        neu_info = neumaticos.get(m['codigo'], {})
        codigos_conteo[m['codigo']].append({
            'placa': m['placa'],
            'posicion': m['posicion'],
            'km': m['km'],
            'marca': neu_info.get('marca', ''),
            'medida': neu_info.get('medida', ''),
            'diseno': neu_info.get('diseno', '')
        })
    duplicados = {k: v for k, v in codigos_conteo.items() if len(v) > 1}

    n_tractos = len(tractos)
    n_teoricos = n_tractos * 24
    n_movimientos = len(movimientos)
    n_faltantes = n_teoricos - n_movimientos
    n_unicos = len(neumaticos)
    n_duplicados = len(duplicados)
    n_apariciones_dup = sum(len(v) for v in duplicados.values())
    diferencia_mov_uni = n_movimientos - n_unicos

    # =========================================================================
    # HOJA 1: RESUMEN_CONCILIACION
    # =========================================================================
    ws1 = wb.create_sheet(title="RESUMEN_CONCILIACION")
    ws1.views.sheetView[0].showGridLines = True

    # Banner de Título
    ws1.merge_cells("A1:G2")
    ws1["A1"] = "INFORME DE AUDITORÍA Y CONCILIACIÓN DE NEUMÁTICOS"
    ws1["A1"].font = font_titulo
    ws1["A1"].fill = fill_navy
    ws1["A1"].alignment = align_center

    ws1.merge_cells("A3:G3")
    ws1["A3"] = f"Transportes Díaz Acarreos Generales S.A.C. | Flota Activa ({n_tractos} Unidades según Resumen!A2:A35)"
    ws1["A3"].font = font_subtitulo
    ws1["A3"].fill = fill_navy
    ws1["A3"].alignment = align_center

    # Sección 1: Cuadro Resumen de Conciliación Numérica
    ws1["A5"] = "1. CONCILIACIÓN MATEMÁTICA DEL SISTEMA"
    ws1["A5"].font = font_seccion

    headers_kpi = ["MÉTRICA / CONCEPTO", "CANTIDAD", "DETALLE TÉCNICO Y JUSTIFICACIÓN"]
    for col_idx, h in enumerate(headers_kpi, start=1):
        c = ws1.cell(row=6, column=col_idx, value=h)
        c.font = font_cabecera
        c.fill = fill_cabecera
        c.alignment = align_center
        c.border = border_header

    filas_kpi = [
        ("Unidades Operativas Activas (Flota)", n_tractos, f"Filtro estricto según pestaña 'Resumen!A2:A35' ({n_tractos} tractocamiones + {len(cisternas)} cisternas vinculadas)"),
        (f"Capacidad Teórica Máxima ({n_tractos} x 24 posiciones)", n_teoricos, "Cada unidad tracto-cisterna tiene un cupo máximo de 24 posiciones (11 en tracto + 13 en cisterna)"),
        ("Posiciones Faltantes / Sin Código en k.xlsx", n_faltantes, f"Llantas de repuesto ausentes (RT/RS) o celdas marcadas 'Sin codigo' en 7 vehículos específicos"),
        ("Total Asignaciones / Montajes Físicos Extraídos", n_movimientos, f"Montajes reales instalados en las {n_tractos} unidades activas ({n_teoricos} teóricos - {n_faltantes} faltantes)"),
        ("Códigos Seriales con Duplicidad en k.xlsx", n_duplicados, f"Códigos que figuran repetidos en 2 posiciones o unidades diferentes (generan {n_apariciones_dup} apariciones físicas)"),
        ("Total Neumáticos Únicos (Catálogo Maestro DAG)", n_unicos, f"Catálogo consolidado de activos físicos bajo la OPCIÓN B ({n_movimientos} montajes - {diferencia_mov_uni} ocurrencias duplicadas)")
    ]

    for idx, (concepto, cant, desc) in enumerate(filas_kpi, start=7):
        c1 = ws1.cell(row=idx, column=1, value=concepto)
        c2 = ws1.cell(row=idx, column=2, value=cant)
        c3 = ws1.cell(row=idx, column=3, value=desc)
        c1.font = font_bold
        c2.font = font_bold
        c3.font = font_normal
        c1.border = border_celda
        c2.border = border_celda
        c3.border = border_celda
        c2.alignment = align_center
        if "Neumáticos Únicos" in concepto:
            c1.fill = fill_kpi
            c2.fill = fill_kpi
            c3.fill = fill_kpi
        elif "Posiciones Faltantes" in concepto:
            c1.fill = fill_alerta
            c2.fill = fill_alerta
            c3.fill = fill_alerta

    # Sección 2: Explicación de la Opción B y Casos Duplicados
    row_exp = 15
    ws1.cell(row=row_exp, column=1, value=f"2. ¿POR QUÉ NO HAY {n_teoricos} Y POR QUÉ SE ADOPTA LA OPCIÓN B ({n_unicos} ÚNICOS)?").font = font_seccion

    explicaciones = [
        (f"A. ¿Por qué no hay {n_teoricos} montajes sino {n_movimientos}?",
         f"Si multiplicamos {n_tractos} vehículos por 24 posiciones (11 en tracto: P1..P10, RT; y 13 en cisterna: P11..P22, RS), el total teórico sería {n_teoricos}.\n"
         "Sin embargo, en la operación real, 7 unidades no tienen su dotación completa en la fecha de inspección:\n"
         "  • 6 unidades no llevan llanta de repuesto en tracto (RT) o en cisterna (RS) en el momento de la inspección.\n"
         "  • La unidad T8U-906 tiene 3 llantas rodantes con la anotación 'Sin codigo' en k.xlsx y carece de repuestos.\n"
         f"Esto suma exactamente {n_faltantes} posiciones no instaladas o sin código: {n_teoricos} - {n_faltantes} = {n_movimientos} montajes físicos reales."),

        (f"B. ¿Por qué {n_unicos} Neumáticos Únicos vs. {n_movimientos} Montajes? (La Opción B)",
         "En una base de datos relacional (como la base de datos de DAG Neumáticos), la tabla 'Neumáticos' exige que cada código de serie sea único (Llave Primaria o Primary Key).\n"
         "No es posible almacenar dos registros con el mismo código como dos neumáticos distintos.\n"
         "Por ello, la Opción B establece:\n"
         f"  1. Plantilla_Registro_Masivo_Neumaticos_DAG.xlsx contiene exactamente los {n_unicos} activos únicos, garantizando integridad referencial.\n"
         f"  2. Plantilla_Asignacion_Movimiento_Neumaticos_DAG.xlsx registra los {n_movimientos} montajes físicos reales, preservando la trazabilidad de dónde está ubicada cada llanta."),

        (f"C. Causa raíz de los {n_duplicados} duplicados en el archivo k.xlsx",
         f"Los {n_duplicados} casos duplicados (que representan {n_apariciones_dup} montajes para {n_duplicados} códigos únicos) se deben a causas operativas en el llenado manual del Excel:\n"
         "  1. Llantas de repuesto no actualizadas: En varios casos, una llanta figura montada en una unidad activa y a la vez anotada como repuesto (RT o RS) en otra.\n"
         "  2. Rotaciones entre unidades sin dar de baja la hoja anterior: El neumático fue transferido a otro camión, pero no se borró de la pestaña origen.\n"
         "  3. Corrección de tipografías: Cuando se corrigen errores de digitación en k.xlsx, el número de duplicados se reduce automáticamente.")
    ]

    r_curr = row_exp + 1
    for titulo_exp, cuerpo_exp in explicaciones:
        ws1.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=7)
        c_tit = ws1.cell(row=r_curr, column=1, value=titulo_exp)
        c_tit.font = font_bold
        c_tit.fill = fill_subcabecera

        ws1.merge_cells(start_row=r_curr + 1, start_column=1, end_row=r_curr + 3, end_column=7)
        c_cue = ws1.cell(row=r_curr + 1, column=1, value=cuerpo_exp)
        c_cue.font = font_normal
        c_cue.alignment = Alignment(wrap_text=True, vertical="top")

        r_curr += 5

    # =========================================================================
    # HOJA 2: MENOS_DE_24_NEUMATICOS
    # =========================================================================
    ws2 = wb.create_sheet(title="MENOS_DE_24_NEUMATICOS")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G2")
    ws2["A1"] = "DETALLE DE VEHÍCULOS CON MENOS DE 24 NEUMÁTICOS (12 FALTANTES)"
    ws2["A1"].font = font_titulo
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = align_center

    ws2.merge_cells("A3:G3")
    ws2["A3"] = "Explicación exacta de por qué el total de montajes es 804 y no 816 (Capacidad: 34 x 24 = 816)"
    ws2["A3"].font = font_subtitulo
    ws2["A3"].fill = fill_navy
    ws2["A3"].alignment = align_center

    headers_incompletas = [
        "TRACTO", "CISTERNA", "INSTALADOS (de 24)", "TRACTO (de 11)", "CISTERNA (de 13)",
        "FALTANTES", "POSICIONES FALTANTES Y MOTIVO EN ARCHIVO MAESTRO (k.xlsx)"
    ]
    for col_idx, h in enumerate(headers_incompletas, start=1):
        c = ws2.cell(row=5, column=col_idx, value=h)
        c.font = font_cabecera
        c.fill = fill_cabecera
        c.alignment = align_center
        c.border = border_header

    # Datos específicos extraídos de k.xlsx
    datos_incompletas = [
        ("ACF-943", "TBC-975", 22, 10, 12, 2, "Falta Tracto RT (Fila 14 vacía) y Cisterna RS (Fila 27 vacía). Sin repuestos."),
        ("AKO-813", "TED-983", 23, 10, 13, 1, "Falta Tracto RT (Fila 14 marcada como 'Sin código' marca Onyx)."),
        ("AMY-787", "TEF-986", 23, 10, 13, 1, "Falta Tracto RT (Fila 14 vacía). Sin repuesto en tracto."),
        ("AUW-702", "THV-976", 23, 11, 12, 1, "Falta Cisterna RS (Fila 27 marcada como 'SIN CÓDIGO' marca ONYX)."),
        ("T5Q-937", "TEE-972", 23, 11, 12, 1, "Falta Cisterna RS (Fila 27 vacía, marca Lander sin serie)."),
        ("T5R-834", "T8E-976", 23, 11, 12, 1, "Falta Cisterna RS (Fila 27 vacía). Sin repuesto en cisterna."),
        ("T8U-906", "TEB-977", 19, 7, 12, 5, "Faltan 5 posiciones: Tracto P3, P7, P8 con texto 'Sin codigo'; Tracto RT con 'Sin codigo'; Cisterna RS vacía.")
    ]

    for idx, (tr, ci, inst, n_tr, n_ci, falt, detalle) in enumerate(datos_incompletas, start=6):
        c_tr = ws2.cell(row=idx, column=1, value=tr)
        c_ci = ws2.cell(row=idx, column=2, value=ci)
        c_in = ws2.cell(row=idx, column=3, value=f"{inst} / 24")
        c_nt = ws2.cell(row=idx, column=4, value=f"{n_tr} / 11")
        c_nc = ws2.cell(row=idx, column=5, value=f"{n_ci} / 13")
        c_fa = ws2.cell(row=idx, column=6, value=falt)
        c_de = ws2.cell(row=idx, column=7, value=detalle)

        for cell in [c_tr, c_ci, c_in, c_nt, c_nc, c_fa]:
            cell.alignment = align_center
            cell.font = font_bold
            cell.border = border_celda
        c_de.alignment = align_left
        c_de.font = font_normal
        c_de.border = border_celda
        c_fa.fill = fill_alerta

    # Fila Totalizadora
    r_tot = 6 + len(datos_incompletas)
    ws2.cell(row=r_tot, column=1, value="TOTAL GENERAL").font = font_bold
    ws2.cell(row=r_tot, column=1).alignment = align_center
    ws2.cell(row=r_tot, column=1).border = border_header
    ws2.cell(row=r_tot, column=6, value=12).font = font_bold
    ws2.cell(row=r_tot, column=6).alignment = align_center
    ws2.cell(row=r_tot, column=6).fill = fill_alerta
    ws2.cell(row=r_tot, column=6).border = border_header
    ws2.cell(row=r_tot, column=7, value="Suma exacta de 12 posiciones faltantes (816 Teóricos - 12 Faltantes = 804 Montajes Reales)").font = font_bold
    ws2.cell(row=r_tot, column=7).border = border_header

    # =========================================================================
    # HOJA 3: DETALLE_DUPLICADOS
    # =========================================================================
    ws3 = wb.create_sheet(title="DETALLE_DUPLICADOS")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:H2")
    ws3["A1"] = "REGISTRO Y DETALLE DE CÓDIGOS DE NEUMÁTICOS DUPLICADOS EN k.xlsx"
    ws3["A1"].font = font_titulo
    ws3["A1"].fill = fill_navy
    ws3["A1"].alignment = align_center

    ws3.merge_cells("A3:H3")
    ws3["A3"] = f"{n_duplicados} códigos seriales repetidos que generan una diferencia de {diferencia_mov_uni} entre Montajes ({n_movimientos}) y Neumáticos Únicos ({n_unicos})"
    ws3["A3"].font = font_subtitulo
    ws3["A3"].fill = fill_navy
    ws3["A3"].alignment = align_center

    headers_dups = [
        "ITEM", "CÓDIGO SERIAL", "OCURRENCIAS", "UBICACIÓN 1 (Placa / Posición / Marca)",
        "UBICACIÓN 2 (Placa / Posición / Marca)", "TIPO DE DUPLICIDAD",
        "ODÓMETRO (Km)", "DIAGNÓSTICO EN ARCHIVO MAESTRO"
    ]
    for col_idx, h in enumerate(headers_dups, start=1):
        c = ws3.cell(row=5, column=col_idx, value=h)
        c.font = font_cabecera
        c.fill = fill_cabecera
        c.alignment = align_center
        c.border = border_header

    r_dup = 6
    for i, (cod, lista) in enumerate(sorted(duplicados.items(), key=lambda x: str(x[0])), start=1):
        u1 = lista[0]
        u2 = lista[1] if len(lista) > 1 else {}

        # Determinar tipo
        es_misma_placa = (u1['placa'] == u2.get('placa'))
        involucra_repuesto = ('R' in u1['posicion'] or 'R' in u2.get('posicion', ''))

        if es_misma_placa:
            tipo = "Mismo Vehículo (Doble Digitación)"
            diag = f"Registrado en posiciones adyacentes de {u1['placa']}. Error de digitación en inspección."
        elif involucra_repuesto:
            tipo = "Repuesto Cruzado"
            diag = "Figura instalado en una unidad y como repuesto en otra. Falta actualizar cambio físico."
        else:
            tipo = "Rotación entre Unidades"
            diag = "Neumático transferido a otra unidad sin haber sido borrado de la hoja anterior."

        c_it = ws3.cell(row=r_dup, column=1, value=i)
        c_co = ws3.cell(row=r_dup, column=2, value=str(cod))
        c_oc = ws3.cell(row=r_dup, column=3, value=len(lista))
        c_u1 = ws3.cell(row=r_dup, column=4, value=f"{u1['placa']} [{u1['posicion']}] - {u1['marca']}")
        c_u2 = ws3.cell(row=r_dup, column=5, value=f"{u2.get('placa')} [{u2.get('posicion')}] - {u2.get('marca')}")
        c_tp = ws3.cell(row=r_dup, column=6, value=tipo)
        c_km = ws3.cell(row=r_dup, column=7, value=f"Km {u1['km']:,.0f}")
        c_dg = ws3.cell(row=r_dup, column=8, value=diag)

        for cell in [c_it, c_co, c_oc, c_tp, c_km]:
            cell.alignment = align_center
            cell.font = font_bold
            cell.border = border_celda
        c_co.fill = fill_subcabecera

        for cell in [c_u1, c_u2, c_dg]:
            cell.alignment = align_left
            cell.font = font_normal
            cell.border = border_celda

        if es_misma_placa:
            c_tp.fill = fill_alerta
        r_dup += 1

    # Fila Resumen Duplicados
    ws3.cell(row=r_dup, column=1, value="TOTAL").font = font_bold
    ws3.cell(row=r_dup, column=1).alignment = align_center
    ws3.cell(row=r_dup, column=1).border = border_header
    ws3.cell(row=r_dup, column=2, value=f"{n_duplicados} Códigos").font = font_bold
    ws3.cell(row=r_dup, column=2).alignment = align_center
    ws3.cell(row=r_dup, column=2).border = border_header
    ws3.cell(row=r_dup, column=3, value=f"{n_apariciones_dup} Asignaciones").font = font_bold
    ws3.cell(row=r_dup, column=3).alignment = align_center
    ws3.cell(row=r_dup, column=3).border = border_header
    ws3.cell(row=r_dup, column=4, value=f"Diferencia exacta: {n_apariciones_dup} - {n_duplicados} = {diferencia_mov_uni} montajes en exceso respecto a los {n_unicos} activos únicos.").font = font_bold
    ws3.cell(row=r_dup, column=4).border = border_header

    # =========================================================================
    # HOJA 4: FLOTA_COMPLETA_34
    # =========================================================================
    ws4 = wb.create_sheet(title="FLOTA_COMPLETA_34")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:H2")
    ws4["A1"] = "CONCILIACIÓN INTEGRAL DE LAS 34 UNIDADES DE LA FLOTA ACTIVA"
    ws4["A1"].font = font_titulo
    ws4["A1"].fill = fill_navy
    ws4["A1"].alignment = align_center

    ws4.merge_cells("A3:H3")
    ws4["A3"] = "Estado de dotación de neumáticos, odómetro y presencia de duplicados por unidad"
    ws4["A3"].font = font_subtitulo
    ws4["A3"].fill = fill_navy
    ws4["A3"].alignment = align_center

    headers_flota = [
        "N°", "PLACA TRACTO", "PLACA CISTERNA", "NEUMÁTICOS TRACTO (Max 11)",
        "NEUMÁTICOS CISTERNA (Max 13)", "TOTAL INSTALADO (Max 24)",
        "ESTADO DOTACIÓN", "¿TIENE CÓDIGOS DUPLICADOS?"
    ]
    for col_idx, h in enumerate(headers_flota, start=1):
        c = ws4.cell(row=5, column=col_idx, value=h)
        c.font = font_cabecera
        c.fill = fill_cabecera
        c.alignment = align_center
        c.border = border_header

    r_flota = 6
    for i, (tr, tr_data) in enumerate(sorted(tractos.items()), start=1):
        ci = tr_data.get('cisterna', 'SIN CISTERNA')
        m_tr = movs_por_unidad[tr]['tracto']
        m_ci = movs_por_unidad[tr]['cisterna']

        n_tr = len(m_tr)
        n_ci = len(m_ci)
        total_u = n_tr + n_ci

        # Buscar si alguno de sus neumaticos es duplicado
        cods_u = [m['codigo'] for m in m_tr + m_ci]
        dups_en_u = [str(c) for c in cods_u if c in duplicados]

        c_num = ws4.cell(row=r_flota, column=1, value=i)
        c_tra = ws4.cell(row=r_flota, column=2, value=tr)
        c_cis = ws4.cell(row=r_flota, column=3, value=ci)
        c_ntr = ws4.cell(row=r_flota, column=4, value=f"{n_tr} / 11")
        c_nci = ws4.cell(row=r_flota, column=5, value=f"{n_ci} / 13")
        c_tot = ws4.cell(row=r_flota, column=6, value=f"{total_u} / 24")

        estado = "COMPLETO (24/24)" if total_u == 24 else f"INCOMPLETO ({total_u}/24, faltan {24-total_u})"
        c_est = ws4.cell(row=r_flota, column=7, value=estado)

        tiene_dup_txt = f"SÍ ({', '.join(dups_en_u)})" if dups_en_u else "NO"
        c_dup = ws4.cell(row=r_flota, column=8, value=tiene_dup_txt)

        for cell in [c_num, c_tra, c_cis, c_ntr, c_nci, c_tot, c_est, c_dup]:
            cell.alignment = align_center
            cell.font = font_normal
            cell.border = border_celda

        c_tra.font = font_bold
        c_cis.font = font_bold
        c_tot.font = font_bold

        if total_u < 24:
            c_tot.fill = fill_alerta
            c_est.fill = fill_alerta
            c_est.font = font_alerta
        else:
            c_est.font = font_bold

        if dups_en_u:
            c_dup.fill = fill_subcabecera
            c_dup.font = font_bold

        r_flota += 1

    # Fila totalizadora flota
    ws4.cell(row=r_flota, column=1, value="TOTALES").font = font_bold
    ws4.cell(row=r_flota, column=1).alignment = align_center
    ws4.cell(row=r_flota, column=1).border = border_header
    ws4.cell(row=r_flota, column=2, value="34 Tractos").font = font_bold
    ws4.cell(row=r_flota, column=2).alignment = align_center
    ws4.cell(row=r_flota, column=2).border = border_header
    ws4.cell(row=r_flota, column=3, value="34 Cisternas").font = font_bold
    ws4.cell(row=r_flota, column=3).alignment = align_center
    ws4.cell(row=r_flota, column=3).border = border_header
    ws4.cell(row=r_flota, column=6, value=f"{len(movimientos)} Montajes").font = font_bold
    ws4.cell(row=r_flota, column=6).alignment = align_center
    ws4.cell(row=r_flota, column=6).fill = fill_kpi
    ws4.cell(row=r_flota, column=6).border = border_header
    ws4.cell(row=r_flota, column=7, value="804 de 816 Teóricos").font = font_bold
    ws4.cell(row=r_flota, column=7).alignment = align_center
    ws4.cell(row=r_flota, column=7).border = border_header

    # Autoajustar anchos de columnas en todas las hojas
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Omitir filas fusionadas para no ensanchar indebidamente
                if cell.row in [1, 2, 3, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25]:
                    continue
                val = str(cell.value or '')
                if val:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Ajustes finos de columnas específicas
    ws1.column_dimensions['A'].width = 38
    ws1.column_dimensions['B'].width = 16
    ws1.column_dimensions['C'].width = 90

    ws2.column_dimensions['A'].width = 16
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 22
    ws2.column_dimensions['D'].width = 18
    ws2.column_dimensions['E'].width = 18
    ws2.column_dimensions['F'].width = 14
    ws2.column_dimensions['G'].width = 95

    ws3.column_dimensions['A'].width = 8
    ws3.column_dimensions['B'].width = 18
    ws3.column_dimensions['C'].width = 14
    ws3.column_dimensions['D'].width = 38
    ws3.column_dimensions['E'].width = 38
    ws3.column_dimensions['F'].width = 32
    ws3.column_dimensions['G'].width = 16
    ws3.column_dimensions['H'].width = 75

    ws4.column_dimensions['A'].width = 8
    ws4.column_dimensions['B'].width = 18
    ws4.column_dimensions['C'].width = 18
    ws4.column_dimensions['D'].width = 24
    ws4.column_dimensions['E'].width = 24
    ws4.column_dimensions['F'].width = 24
    ws4.column_dimensions['G'].width = 34
    ws4.column_dimensions['H'].width = 34

    # Crear carpeta de destino si no existe
    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    wb.save(ruta_salida)
    return ruta_salida


if __name__ == '__main__':
    import sys
    base_local = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    if base_local not in sys.path:
        sys.path.insert(0, base_local)
    from scripts.etl_maestro import extraer_datos_maestro
    maestro_path = os.path.join(base_local, 'Archivos', 'Maestro', 'k.xlsx')
    datos = extraer_datos_maestro(maestro_path)
    
    # Guardar en Plantillas y en Reportes
    salida1 = os.path.join(base_local, 'Archivos', 'Plantillas', 'Reporte_Auditoria_Conciliacion_Neumaticos_DAG.xlsx')
    salida2 = os.path.join(base_local, 'Archivos', 'Reportes', 'Reporte_Auditoria_Conciliacion_Neumaticos_DAG.xlsx')
    
    crear_reporte_auditoria_excel(datos, salida1)
    crear_reporte_auditoria_excel(datos, salida2)
    print(f"Reporte generado exitosamente en:\n - {salida1}\n - {salida2}")
