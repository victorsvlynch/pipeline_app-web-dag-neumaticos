"""
ETL Pipeline: Extracción de datos desde el archivo Maestro (k.xlsx)
filtrando exclusivamente las placas de la pestaña 'Resumen' (celdas A2 a A35)
y distribución hacia las Plantillas de carga masiva de la aplicación DAG Neumáticos.

Transportes Díaz Acarreos Generales S.A.C.
"""

import os
import re
import sys
import time
import logging
from datetime import datetime
import openpyxl

# Configuración de Logging
logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] %(levelname)s: %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger("ETL_MAESTRO_NEUMATICOS")

# Expresiones regulares
REGEX_CISTERNA = re.compile(r'([A-Z0-9]{3}-[A-Z0-9]{3})')
REGEX_PLACA_TOLERANTE = re.compile(r'([A-Z0-9]{3})[- ]([A-Z0-9]{3})')
REGEX_DNI = re.compile(r'\d{8}')

# Posiciones según guía oficial del sistema
POSICIONES_TRACTO = ['P1', 'P2', 'P3', 'P4', 'P5', 'P6', 'P7', 'P8', 'P9', 'P10', 'RT']
POSICIONES_CISTERNA = ['P11', 'P12', 'P13', 'P14', 'P15', 'P16', 'P17', 'P18', 'P19', 'P20', 'P21', 'P22', 'RS']

MAPA_MARCAS = {
    'VOLVO': 'Volvo',
    'INTERNATIONAL': 'International',
    'INTERNACIONAL': 'International',
    'FREIGHTLINER': 'Freightliner',
    'MACK': 'Mack',
    'SCANIA': 'Scania',
    'MERCEDES': 'Mercedes',
    'KENWORTH': 'Kenworth',
    'MAN': 'Man',
    'IVECO': 'Iveco'
}


def es_codigo_valido(code_raw):
    """
    Valida si un valor corresponde a un código de neumático real.
    Descarta valores nulos, textos de marcador de posición ('sin código', 'no tiene', etc.)
    y encabezados accidentales.
    """
    if code_raw is None:
        return False
    code_str = str(code_raw).strip()
    if not code_str:
        return False
    c_lower = code_str.lower()
    descartes = {
        'none', 'null', 'nan', '0', 'n.a', 'n/a', 'na', 'no tiene',
        'sin codigo', 'sin código', 'sin cdigo', 's/c', 's/n',
        'codigo', 'código', 'placa', 'inicio'
    }
    if c_lower in descartes:
        return False
    if re.search(r'^(sin\s+c[oó]digo|no\s+tiene|s/[cn]|sin\s+n[uú]mero)', c_lower):
        return False
    return True


def normalizar_codigo_neumatico(code_raw):
    """Convierte el código a número entero si es numérico, o cadena limpia."""
    code_str = str(code_raw).strip()
    try:
        f_val = float(code_str)
        if f_val.is_integer():
            return int(f_val)
        return int(f_val)
    except (ValueError, TypeError):
        return code_str


def obtener_unidades_activas_resumen(wb, rango_resumen="A2:A35"):
    """
    Extrae las unidades activas de la pestaña 'Resumen' según el rango especificado (por defecto A2:A35).
    Retorna un diccionario: {tracto_placa: cisterna_placa}
    """
    if 'Resumen' not in wb.sheetnames:
        raise ValueError("No se encontró la pestaña 'Resumen' en el archivo maestro.")

    ws_resumen = wb['Resumen']
    unidades = {}

    start_row, end_row = 2, 35
    if rango_resumen:
        m = re.search(r'(\d+):[A-Za-z]*(\d+)', str(rango_resumen))
        if m:
            start_row = int(m.group(1))
            end_row = int(m.group(2))

    for i, r in enumerate(ws_resumen.iter_rows(values_only=True)):
        fila_num = i + 1
        if fila_num < start_row:
            continue
        if fila_num > end_row:
            break

        val = r[0] if r else None
        if val:
            val_str = str(val).strip()
            parts = val_str.split('/')
            tracto = parts[0].strip().upper()
            cisterna = parts[1].strip().upper() if len(parts) > 1 else None

            # Normalizar formato XXX-XXX
            m_tr = REGEX_PLACA_TOLERANTE.search(tracto)
            clean_tr = f"{m_tr.group(1)}-{m_tr.group(2)}" if m_tr else tracto

            clean_ci = None
            if cisterna:
                m_ci = REGEX_PLACA_TOLERANTE.search(cisterna)
                clean_ci = f"{m_ci.group(1)}-{m_ci.group(2)}" if m_ci else cisterna

            unidades[clean_tr] = clean_ci

    logger.info(f"Unidades activas identificadas en 'Resumen!{rango_resumen}': {len(unidades)} unidades.")
    return unidades


def extraer_datos_maestro(ruta_maestro, rango_resumen="A2:A35"):
    """
    Lee el archivo maestro k.xlsx y extrae datos EXCLUSIVAMENTE para las unidades
    indicadas en 'Resumen' según el rango especificado (por defecto A2 a A35).
    """
    logger.info(f"Cargando archivo maestro: {ruta_maestro}")
    if not os.path.exists(ruta_maestro):
        raise FileNotFoundError(f"No se encontró el archivo maestro en: {ruta_maestro}")

    wb = openpyxl.load_workbook(ruta_maestro, data_only=True, read_only=True)

    # 1. Obtener lista estricta de unidades desde Resumen
    unidades_activas = obtener_unidades_activas_resumen(wb, rango_resumen=rango_resumen)
    placas_tracto_validas = set(unidades_activas.keys())

    # 2. Extraer catálogo de conductores desde la hoja 'LISTA'
    conductores_maestro = {}    # clean_tracto -> {nombre, lic, dni}
    todos_conductores = {}      # dni -> {dni, nombre, estado}
    conductores_sin_dni = []

    if 'LISTA' in wb.sheetnames:
        ws_lista = wb['LISTA']
        for i, row in enumerate(ws_lista.iter_rows(values_only=True)):
            if i == 0 or not row or not any(row):
                continue
            nombre = str(row[3] or row[1] or '').strip().upper()
            lic = str(row[2] or '').strip()
            tracto_raw = str(row[6] or '').strip().upper() if len(row) > 6 else ''

            m_tr = REGEX_PLACA_TOLERANTE.search(tracto_raw)
            clean_tr = f"{m_tr.group(1)}-{m_tr.group(2)}" if m_tr else None

            dni_match = REGEX_DNI.search(lic)
            dni = dni_match.group(0) if dni_match else None

            if clean_tr:
                conductores_maestro[clean_tr] = {
                    'nombre': nombre,
                    'dni': dni,
                    'item': i
                }

    # Caso especial T5Q-937: conductor en hoja CONDUCT. y extracción de mochilas
    mochilas_maestro = {}
    if 'CONDUCT.' in wb.sheetnames:
        ws_cond = wb['CONDUCT.']
        for i, row in enumerate(ws_cond.iter_rows(values_only=True)):
            if not row:
                continue
            if len(row) > 3 and 'T5Q-937' in str(row[2]) and 'T5Q-937' not in conductores_maestro:
                conductores_maestro['T5Q-937'] = {
                    'nombre': 'SAAVEDRA ODAR SEGUNDO ANGEL',
                    'dni': '42706482',
                    'item': 29
                }
            if len(row) > 4 and row[2]:
                m_tr = REGEX_PLACA_TOLERANTE.search(str(row[2]))
                if m_tr:
                    clean_c_tr = f"{m_tr.group(1)}-{m_tr.group(2)}"
                    m_val = str(row[4]).strip().upper() if row[4] else ''
                    mochilas_maestro[clean_c_tr] = 50 if 'SI' in m_val else 0

    # 3. Extraer vehículos y neumáticos para las unidades activas
    tractos = {}
    cisternas = {}
    neumaticos = {}
    movimientos = []
    conductores_finales = {}
    asignaciones_finales = {}

    def safe_float(v, default=12.0):
        try:
            if v is not None:
                return float(v)
        except:
            pass
        return default

    def safe_psi(v, default=120):
        try:
            if v is not None:
                f_val = float(v)
                if 50 <= f_val <= 150:
                    return int(f_val)
        except:
            pass
        return default

    def mapear_condicion(cond_raw):
        if cond_raw is None:
            return 'USADO'
        c_str = str(cond_raw).strip().upper()
        if c_str in ('0', 'N', 'NUEVO', 'NUEVA', 'NUEVAS'):
            return 'NUEVO'
        if c_str in ('DIBUJADA', 'DIBUJADO', 'DIBUJADAS'):
            return 'DIBUJADO'
        if any(c_str.startswith(p) for p in ('1R', '2R', '3R', '4R', 'R', '01R', '02R')) or c_str in ('1', '2', '3', '4', 'REENCAUCHE', 'REENCAUCHADO'):
            return 'REENCAUCHADO'
        return 'USADO'

    dni_sintetico_base = 90000000

    for tracto_placa, cisterna_defecto in unidades_activas.items():
        if tracto_placa not in wb.sheetnames:
            logger.warning(f"La hoja para el tracto {tracto_placa} no existe en k.xlsx")
            continue

        ws = wb[tracto_placa]
        rows = []
        for r_idx, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(r)
            if r_idx >= 40:
                break

        if not rows:
            continue

        header_idx = None
        conductor_hoja = None

        for idx, r in enumerate(rows[:10]):
            if r and any(r):
                for c_idx, cell in enumerate(r):
                    if cell and 'CONDUCTOR' in str(cell).upper():
                        for c_val in r[c_idx + 1:]:
                            if c_val and 'CONDUCTOR' not in str(c_val).upper() and 'RUTA' not in str(c_val).upper():
                                conductor_hoja = str(c_val).strip()
                                break
            if r and len(r) > 1 and r[0] == 'MARCA' and r[1] == 'Placa':
                header_idx = idx
                break

        if header_idx is None or header_idx + 2 >= len(rows):
            continue

        r1 = rows[header_idx + 2]
        raw_marca = str(r1[0]).strip() if r1[0] else 'VOLVO'
        marca = MAPA_MARCAS.get(raw_marca.upper(), 'Volvo')

        km_raw = r1[11] if len(r1) > 11 else None
        try:
            km = float(km_raw) if km_raw is not None else 0.0
        except:
            km = 0.0

        fecha_insp_raw = r1[10] if len(r1) > 10 else None
        fecha_insp = None
        if isinstance(fecha_insp_raw, datetime):
            fecha_insp = fecha_insp_raw.strftime('%Y-%m-%d')
        elif fecha_insp_raw:
            fecha_insp = str(fecha_insp_raw).split(' ')[0]

        # Identificar el límite del primer recuadro (el recuadro más reciente)
        # para garantizar que NUNCA se extraigan datos de recuadros históricos inferiores (bloque 2 o posterior)
        limite_primer_recuadro = len(rows)
        for idx_h, r_h in enumerate(rows):
            if idx_h > (header_idx + 2) and r_h and len(r_h) > 1:
                # Si encontramos el inicio de un segundo recuadro (nuevo CONDUCTOR: o nuevo encabezado MARCA/Placa)
                if (r_h[0] == 'MARCA' and r_h[1] == 'Placa') or (r_h[0] and 'CONDUCTOR' in str(r_h[0]).upper()):
                    limite_primer_recuadro = idx_h
                    break

        # Identificar Cisterna de la unidad dentro del primer recuadro
        cisterna_placa = cisterna_defecto
        cisterna_row_idx = None

        for idx_c, r_check in enumerate(rows[header_idx + 12: min(limite_primer_recuadro, len(rows))]):
            if r_check and len(r_check) > 1 and r_check[1]:
                m = REGEX_CISTERNA.search(str(r_check[1]))
                if m:
                    cisterna_placa = m.group(1)
                    cisterna_row_idx = header_idx + 12 + idx_c
                    break

        # Tipo de Aro: detectar desde columna 19 si existe o clasificar por marca
        raw_aro = str(r1[19]).strip() if len(r1) > 19 and r1[19] else ''
        if 'EUROPEO' in raw_aro.upper():
            tipo_aro = 'Europeo'
        elif 'AMERICANO' in raw_aro.upper():
            tipo_aro = 'Americano'
        else:
            tipo_aro = 'Europeo' if marca in ['Volvo', 'Scania', 'Mercedes'] else 'Americano'

        # Capacidad de tanques de combustible (Galones)
        t_izq = 150 if marca in ['Volvo', 'Scania', 'Mercedes'] else 120
        t_der = 150 if marca in ['Volvo', 'Scania', 'Mercedes'] else 120
        t_mochila = mochilas_maestro.get(tracto_placa, 0)

        tractos[tracto_placa] = {
            'placa': tracto_placa,
            'tipo': 'Tractocamion',
            'marca': marca,
            'modelo': 'FM1',
            'tipo_aro': tipo_aro,
            'estado': 'Operativo',
            'km': km,
            'carroceria': '',
            'tanque_lateral_izquierdo': t_izq,
            'tanque_lateral_derecho': t_der,
            'tanque_mochila_auxiliar': t_mochila,
            'cisterna': cisterna_placa,
            'fecha_inspeccion': fecha_insp
        }

        if cisterna_placa:
            cisternas[cisterna_placa] = {
                'placa': cisterna_placa,
                'tipo': 'Semirremolque',
                'marca': 'Tragsac',
                'modelo': 'FM1',
                'tipo_aro': 'Americano',
                'estado': 'Operativo',
                'km': 0,
                'carroceria': 'Cisterna',
                'tanque_lateral_izquierdo': 0,
                'tanque_lateral_derecho': 0,
                'tanque_mochila_auxiliar': 0
            }

        # Conductor de la unidad
        c_info = conductores_maestro.get(tracto_placa)
        if c_info:
            c_nombre = c_info['nombre']
            c_dni = c_info['dni'] if c_info['dni'] else str(dni_sintetico_base + c_info['item'])
        elif conductor_hoja:
            c_nombre = conductor_hoja.upper()
            c_dni = str(dni_sintetico_base + len(conductores_finales) + 1)
        else:
            c_nombre = 'CONDUCTOR NO ASIGNADO'
            c_dni = str(dni_sintetico_base + len(conductores_finales) + 1)

        conductores_finales[c_dni] = {
            'dni': c_dni,
            'nombre': c_nombre,
            'estado': 'ACTIVO'
        }
        asignaciones_finales[tracto_placa] = {
            'placa': tracto_placa,
            'dni': c_dni,
            'nombre': c_nombre
        }

        # Neumáticos del Tractocamión (P1 a P10, RT) estrictamente dentro del primer recuadro
        for p_i, pos in enumerate(POSICIONES_TRACTO):
            row_target = header_idx + 2 + p_i
            if row_target < limite_primer_recuadro and row_target < len(rows):
                tr_row = rows[row_target]
                code_raw = tr_row[2] if len(tr_row) > 2 else None
                if es_codigo_valido(code_raw):
                    code_val = normalizar_codigo_neumatico(code_raw)

                    t_medida = str(tr_row[3]).strip() if len(tr_row) > 3 and tr_row[3] else ('295/80R22.5' if pos in ['P1', 'P2'] else '11R22.5')
                    t_marca = str(tr_row[4]).strip().title() if len(tr_row) > 4 and tr_row[4] else 'Lander'
                    t_tipo = 'DIRECCION' if pos in ['P1', 'P2'] else 'TRACCION'
                    t_diseno = str(tr_row[6]).strip() if len(tr_row) > 6 and tr_row[6] else 'LS215'
                    t_condicion = mapear_condicion(tr_row[5] if len(tr_row) > 5 else None)
                    t_psi = safe_psi(tr_row[16] if len(tr_row) > 16 else None)

                    prof_int = safe_float(tr_row[12] if len(tr_row) > 12 else None)
                    prof_cent = safe_float(tr_row[13] if len(tr_row) > 13 else None)
                    prof_ext = safe_float(tr_row[14] if len(tr_row) > 14 else None)

                    neumaticos[code_val] = {
                        'codigo': code_val,
                        'marca': t_marca,
                        'medida': t_medida,
                        'tipo': t_tipo,
                        'diseno': t_diseno,
                        'condicion': t_condicion,
                        'prof_int': prof_int,
                        'prof_cent': prof_cent,
                        'prof_ext': prof_ext,
                        'semana_dot': 24,
                        'ano_dot': 2023,
                        'psi': t_psi
                    }

                    movimientos.append({
                        'placa': tracto_placa,
                        'codigo': code_val,
                        'posicion': pos,
                        'km': km,
                        'motivo': 'Agregado Masivo Manual'
                    })

        # Neumáticos de la Cisterna (P11 a P22, RS) estrictamente dentro del primer recuadro
        if cisterna_placa and cisterna_row_idx:
            for p_i, pos in enumerate(POSICIONES_CISTERNA):
                row_target = cisterna_row_idx + p_i
                if row_target < limite_primer_recuadro and row_target < len(rows):
                    ci_row = rows[row_target]
                    code_raw = ci_row[2] if len(ci_row) > 2 else None
                    if es_codigo_valido(code_raw):
                        code_val = normalizar_codigo_neumatico(code_raw)

                        t_medida = str(ci_row[3]).strip() if len(ci_row) > 3 and ci_row[3] else '11R22.5'
                        t_marca = str(ci_row[4]).strip().title() if len(ci_row) > 4 and ci_row[4] else 'Lander'
                        t_diseno = str(ci_row[6]).strip() if len(ci_row) > 6 and ci_row[6] else 'LS168'
                        t_condicion = mapear_condicion(ci_row[5] if len(ci_row) > 5 else None)
                        t_psi = safe_psi(ci_row[16] if len(ci_row) > 16 else None)

                        prof_int = safe_float(ci_row[12] if len(ci_row) > 12 else None)
                        prof_cent = safe_float(ci_row[13] if len(ci_row) > 13 else None)
                        prof_ext = safe_float(ci_row[14] if len(ci_row) > 14 else None)

                        neumaticos[code_val] = {
                            'codigo': code_val,
                            'marca': t_marca,
                            'medida': t_medida,
                            'tipo': 'ARRASTRE',
                            'diseno': t_diseno,
                            'condicion': t_condicion,
                            'prof_int': prof_int,
                            'prof_cent': prof_cent,
                            'prof_ext': prof_ext,
                            'semana_dot': 24,
                            'ano_dot': 2023,
                            'psi': t_psi
                        }

                        movimientos.append({
                            'placa': cisterna_placa,
                            'codigo': code_val,
                            'posicion': pos,
                            'km': km,
                            'motivo': 'Agregado Masivo Manual'
                        })

    wb.close()

    logger.info(f"=== EXTRACCIÓN DE UNIDADES ACTIVAS (RESUMEN!{rango_resumen}) COMPLETADA ===")
    logger.info(f" - Tractocamiones: {len(tractos)}")
    logger.info(f" - Semirremolques: {len(cisternas)}")
    logger.info(f" - Conductores: {len(conductores_finales)}")
    logger.info(f" - Neumáticos únicos: {len(neumaticos)}")
    logger.info(f" - Movimientos / Asignaciones: {len(movimientos)}")

    return {
        'conductores': conductores_finales,
        'asignaciones_conductores': asignaciones_finales,
        'tractos': tractos,
        'cisternas': cisternas,
        'neumaticos': neumaticos,
        'movimientos': movimientos,
        'unidades_activas': unidades_activas
    }


def _actualizar_plantilla_excel(ruta_archivo, nombre_hoja, filas_datos, reintentos=3, pausa=0.6):
    """
    Actualiza la hoja de datos principal del archivo plantilla conservando encabezados
    y guías. Incluye reintentos ante bloqueos temporales de archivos en Windows.
    """
    if not os.path.exists(ruta_archivo):
        raise FileNotFoundError(f"Plantilla no encontrada: {ruta_archivo}")

    wb = openpyxl.load_workbook(ruta_archivo)
    try:
        if nombre_hoja not in wb.sheetnames:
            ws = wb.sheetnames[0]
            logger.warning(f"Hoja '{nombre_hoja}' no encontrada, utilizando primera hoja: '{ws.title}'")
        else:
            ws = wb[nombre_hoja]

        # Eliminar datos anteriores conservando la fila 1 (encabezados)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

        for fila in filas_datos:
            ws.append(fila)

        guardado = False
        ultimo_error = None
        for intento in range(1, reintentos + 1):
            try:
                wb.save(ruta_archivo)
                guardado = True
                break
            except PermissionError as e:
                ultimo_error = e
                logger.warning(f"Intento {intento}/{reintentos}: Archivo '{os.path.basename(ruta_archivo)}' bloqueado. Reintentando en {pausa}s...")
                time.sleep(pausa)

        if not guardado:
            msg = (f"No se pudo guardar '{os.path.basename(ruta_archivo)}'. "
                   f"El archivo está bloqueado por otro proceso o abierto en Excel. "
                   f"Por favor ciérrelo e intente nuevamente.")
            logger.error(msg)
            raise PermissionError(msg) from ultimo_error

        logger.info(f"Actualizada plantilla '{os.path.basename(ruta_archivo)}' con {len(filas_datos)} registros.")
    finally:
        wb.close()


def generar_plantilla_vehiculos(plantillas_dir, tractos, cisternas):
    """Genera Plantilla_Registro_Masivo_Vehiculos_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Vehiculos_DAG.xlsx')
    filas = []

    # 1. Tractocamiones
    for p in sorted(tractos.keys()):
        v = tractos[p]
        filas.append([
            v['placa'],
            v['tipo'],
            v['marca'],
            v['modelo'],
            v['tipo_aro'],
            v['estado'],
            v['km'],
            v['carroceria'],
            v.get('tanque_lateral_izquierdo', 0),
            v.get('tanque_lateral_derecho', 0),
            v.get('tanque_mochila_auxiliar', 0)
        ])

    # 2. Semirremolques
    for p in sorted(cisternas.keys()):
        c = cisternas[p]
        filas.append([
            c['placa'],
            c['tipo'],
            c['marca'],
            c['modelo'],
            c['tipo_aro'],
            c['estado'],
            c['km'],
            c['carroceria'],
            c.get('tanque_lateral_izquierdo', 0),
            c.get('tanque_lateral_derecho', 0),
            c.get('tanque_mochila_auxiliar', 0)
        ])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_VEHICULOS', filas)
    return len(filas)


def generar_plantilla_conductores(plantillas_dir, conductores):
    """Genera Plantilla_Registro_Masivo_Conductores_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Conductores_DAG.xlsx')
    filas = []

    for dni in sorted(conductores.keys()):
        c = conductores[dni]
        filas.append([
            c['dni'],
            c['nombre'],
            c['estado']
        ])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_CONDUCTORES', filas)
    return len(filas)


def generar_plantilla_asignacion_conductores(plantillas_dir, asignaciones, tractos, conductores):
    """Genera Plantilla_Registro_Masivo_Asignacion_Conductores_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Asignacion_Conductores_DAG.xlsx')
    filas = []
    fecha_actual = datetime.now().strftime('%Y-%m-%d')

    for placa in sorted(tractos.keys()):
        if placa in asignaciones:
            asig = asignaciones[placa]
            filas.append([
                placa,
                asig['dni'],
                asig['nombre'],
                fecha_actual,
                'ASIGNACIÓN MASIVA AUTOMÁTICA DAG'
            ])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_ASIGNACIONES', filas)
    return len(filas)


def generar_plantilla_neumaticos(plantillas_dir, neumaticos):
    """Genera Plantilla_Registro_Masivo_Neumaticos_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Neumaticos_DAG.xlsx')
    filas = []

    for codigo in sorted(neumaticos.keys(), key=lambda x: str(x)):
        n = neumaticos[codigo]
        filas.append([
            n['codigo'],
            n['marca'],
            n['medida'],
            n['tipo'],
            n['diseno'],
            n['condicion'],
            n['prof_int'],
            n['prof_cent'],
            n['prof_ext'],
            n['semana_dot'],
            n['ano_dot'],
            n['psi']
        ])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_NEUMATICOS', filas)
    return len(filas)


def generar_plantilla_movimientos(plantillas_dir, movimientos):
    """Genera Plantilla_Asignacion_Movimiento_Neumaticos_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Asignacion_Movimiento_Neumaticos_DAG.xlsx')
    filas = []

    for m in movimientos:
        filas.append([
            m['placa'],
            m['codigo'],
            m['posicion'],
            m['km'],
            m['motivo']
        ])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_MOVIMIENTOS', filas)
    return len(filas)


def generar_plantilla_comparativa(plantillas_dir, tractos):
    """Genera Plantilla_Registro_Masivo_Comparativa_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Comparativa_DAG.xlsx')
    filas = []

    for placa in sorted(tractos.keys()):
        t = tractos[placa]
        cap_total = t.get('tanque_lateral_izquierdo', 0) + t.get('tanque_lateral_derecho', 0) + t.get('tanque_mochila_auxiliar', 0)
        galones_diesel = cap_total if cap_total > 0 else (223 if t.get('marca') == 'Volvo' else 200)

        filas.append([placa, 'GASOLINA REGULAR', None, 'ACTIVO'])
        filas.append([placa, 'GASOLINA PREMIUM', None, 'ACTIVO'])
        filas.append([placa, 'DIESEL B5 S-50', galones_diesel, 'ACTIVO'])

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_COMPARATIVA', filas)
    return len(filas)


def generar_plantilla_grifos(plantillas_dir):
    """Genera / Asegura Plantilla_Registro_Masivo_Grifos_DAG.xlsx"""
    archivo = os.path.join(plantillas_dir, 'Plantilla_Registro_Masivo_Grifos_DAG.xlsx')
    grifos_base = [
        ['GRIFO PRIMAX PIURA CENTRO', 'PRIMAX', 'PIURA', 'AV. SANCHEZ CERRO 1234', 'ACTIVO'],
        ['ESTACION REPSOL TALARA SUR', 'REPSOL', 'TALARA', 'PANAMERICANA NORTE KM 1080', 'ACTIVO'],
        ['PATIO CENTRAL DAG PAITA', 'PATIO / PROPIO', 'PAITA', 'ZONA INDUSTRIAL MZ. C LOTE 4', 'ACTIVO'],
        ['GRIFO PETROPERU CHICLAYO', 'PETROPERÚ', 'CHICLAYO', 'AV. VICTOR RAUL HAYA DE LA TORRE 500', 'ACTIVO'],
        ['GRIFO SHELL LIMA NORTE', 'SHELL', 'LIMA', 'PANAMERICANA NORTE KM 22', 'ACTIVO']
    ]

    _actualizar_plantilla_excel(archivo, 'PLANTILLA_GRIFOS', grifos_base)
    return len(grifos_base)


def ejecutar_pipeline(maestro_path=None, plantillas_dir=None, rango_resumen="A2:A35"):
    """
    Función principal de ejecución del Pipeline ETL.
    """
    inicio = datetime.now()
    logger.info(f"=== INICIANDO PIPELINE ETL (FILTRO: RESUMEN!{rango_resumen}) ===")

    base_local = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

    if not maestro_path:
        local_maestro = os.path.join(base_local, 'Archivos', 'Maestro', 'k.xlsx')
        if os.path.exists(local_maestro):
            maestro_path = local_maestro
        elif os.path.exists('/app/Archivos/Maestro/k.xlsx'):
            maestro_path = '/app/Archivos/Maestro/k.xlsx'
        else:
            maestro_path = local_maestro

    if not plantillas_dir:
        local_plantillas = os.path.join(base_local, 'Archivos', 'Plantillas')
        if os.path.exists(local_plantillas):
            plantillas_dir = local_plantillas
        elif os.path.exists('/app/Archivos/Plantillas'):
            plantillas_dir = '/app/Archivos/Plantillas'
        else:
            plantillas_dir = local_plantillas

    logger.info(f"Ruta Maestro: {maestro_path}")
    logger.info(f"Directorio Plantillas: {plantillas_dir}")

    # 1. Extracción con filtro Resumen
    datos = extraer_datos_maestro(maestro_path, rango_resumen=rango_resumen)

    # 2. Generación de las 7 plantillas
    n_vehiculos = generar_plantilla_vehiculos(plantillas_dir, datos['tractos'], datos['cisternas'])
    n_conductores = generar_plantilla_conductores(plantillas_dir, datos['conductores'])
    n_asig_conductores = generar_plantilla_asignacion_conductores(plantillas_dir, datos['asignaciones_conductores'], datos['tractos'], datos['conductores'])
    n_neumaticos = generar_plantilla_neumaticos(plantillas_dir, datos['neumaticos'])
    n_movimientos = generar_plantilla_movimientos(plantillas_dir, datos['movimientos'])
    n_comparativa = generar_plantilla_comparativa(plantillas_dir, datos['tractos'])
    n_grifos = generar_plantilla_grifos(plantillas_dir)

    # 3. Generación del Reporte de Auditoría y Conciliación (Opción B, Duplicados y Faltantes)
    reporte_auditoria_path = os.path.join(plantillas_dir, 'Reporte_Auditoria_Conciliacion_Neumaticos_DAG.xlsx')
    try:
        try:
            from scripts.generar_reporte_auditoria import crear_reporte_auditoria_excel
        except ModuleNotFoundError:
            from generar_reporte_auditoria import crear_reporte_auditoria_excel
        crear_reporte_auditoria_excel(datos, reporte_auditoria_path)
        dir_reportes = os.path.join(base_local, 'Archivos', 'Reportes')
        if os.path.exists(dir_reportes):
            crear_reporte_auditoria_excel(datos, os.path.join(dir_reportes, 'Reporte_Auditoria_Conciliacion_Neumaticos_DAG.xlsx'))
        logger.info(f"Reporte de Auditoría y Conciliación generado en: {reporte_auditoria_path}")
    except Exception as e_rep:
        logger.warning(f"No se pudo generar el reporte de auditoría automático: {e_rep}")

    duracion = (datetime.now() - inicio).total_seconds()
    resumen = {
        'duracion_segundos': duracion,
        'vehiculos_total': n_vehiculos,
        'conductores_total': n_conductores,
        'asignaciones_conductores_total': n_asig_conductores,
        'neumaticos_total': n_neumaticos,
        'movimientos_total': n_movimientos,
        'comparativas_total': n_comparativa,
        'grifos_total': n_grifos,
        'reporte_auditoria': reporte_auditoria_path
    }

    logger.info("=== PIPELINE ETL COMPLETADO EXITOSAMENTE ===")
    logger.info("Resumen de registros generados en plantillas:")
    for k, v in resumen.items():
        logger.info(f" - {k}: {v}")

    return resumen


if __name__ == '__main__':
    ejecutar_pipeline()
